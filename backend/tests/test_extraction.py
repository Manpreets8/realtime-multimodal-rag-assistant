import zipfile
from datetime import date, datetime
from pathlib import Path

import docx
import pymupdf
import pytest

from app.rag import extraction
from app.rag.extraction import ExtractionError, extract_text


def make_pdf(path: Path, pages: list[str], **save_options) -> Path:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_textbox(pymupdf.Rect(72, 72, 540, 720), text, fontsize=11)
    document.save(path, **save_options)
    return path


def test_pdf_text_is_extracted_per_page(tmp_path: Path) -> None:
    path = make_pdf(
        tmp_path / "handbook.pdf",
        ["Annual leave is 18 days.\n\nSick leave is 10 days.", "Remote work requires manager approval."],
    )

    result = extract_text(path, ".pdf")

    assert result.page_count == 2
    assert [block.page_number for block in result.blocks] == [1, 2]
    assert "Annual leave is 18 days." in result.blocks[0].text
    assert "Remote work requires manager approval." in result.blocks[1].text


def test_pdf_hyphenated_line_breaks_are_rejoined(tmp_path: Path) -> None:
    path = make_pdf(tmp_path / "hyphen.pdf", ["The employ-\nment contract is binding."])

    assert "employment contract" in extract_text(path, ".pdf").blocks[0].text


def test_pdf_pages_without_text_are_skipped_but_counted(tmp_path: Path) -> None:
    path = make_pdf(tmp_path / "scan.pdf", ["", "Only page two has text."])

    result = extract_text(path, ".pdf")

    assert result.page_count == 2
    assert [block.page_number for block in result.blocks] == [2]


def test_encrypted_pdf_is_rejected_with_a_clear_message(tmp_path: Path) -> None:
    path = make_pdf(
        tmp_path / "secret.pdf",
        ["Confidential"],
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-pass",
        user_pw="user-pass",
    )

    with pytest.raises(ExtractionError, match="password-protected"):
        extract_text(path, ".pdf")


def test_corrupt_pdf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\n this is not really a pdf body")

    with pytest.raises(ExtractionError, match="damaged"):
        extract_text(path, ".pdf")


def test_docx_sections_follow_headings_and_include_tables(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Welcome to the company.")
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees receive 18 days of annual leave.")
    document.add_heading("Carry over", level=2)
    document.add_paragraph("Up to 5 unused days carry over.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Grade", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Senior", "22"
    document.add_heading("Benefits", level=1)
    document.add_paragraph("Health insurance is provided.")
    path = tmp_path / "policy.docx"
    document.save(path)

    blocks = extract_text(path, ".docx").blocks

    assert [block.section for block in blocks] == [
        None,
        "Leave Policy",
        "Leave Policy > Carry over",
        "Benefits",
    ]
    assert blocks[1].text.startswith("Leave Policy")
    assert "Grade | Days\nSenior | 22" in blocks[2].text
    assert all(block.page_number is None for block in blocks)


def test_docx_zip_bomb_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extraction, "MAX_DOCX_COMPRESSION_RATIO", 50)
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 2_000_000)  # compresses ~1000:1

    with pytest.raises(ExtractionError, match="too large when decompressed"):
        extract_text(path, ".docx")


def test_markdown_sections_ignore_headings_inside_code_blocks(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text(
        "Intro text.\n\n# Setup\nInstall it.\n\n```bash\n# not a heading\npip install app\n```\n\n"
        "## Configure\nSet the env vars.\n\n# Usage\nRun it.\n",
        encoding="utf-8",
    )

    blocks = extract_text(path, ".md").blocks

    assert [block.section for block in blocks] == [None, "Setup", "Setup > Configure", "Usage"]
    assert "# not a heading" in blocks[1].text


def test_utf16_text_file(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes("Café menu — espresso".encode("utf-16"))

    blocks = extract_text(path, ".txt").blocks

    assert blocks[0].text == "Café menu — espresso"
    assert blocks[0].page_number is None


def test_missing_file_is_reported(tmp_path: Path) -> None:
    with pytest.raises(ExtractionError, match="missing"):
        extract_text(tmp_path / "gone.pdf", ".pdf")


# --- metadata -----------------------------------------------------------------------------


def test_pdf_metadata_is_read_and_cleaned(tmp_path: Path) -> None:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Quarterly results.")
    document.set_metadata(
        {
            "title": "Microsoft Word - Q3\u200b Report",  # converter prefix and a zero-width space
            "author": "  Asha   Verma ",
            "creationDate": "D:20240115093000+01'00'",
        }
    )
    path = tmp_path / "report.pdf"
    document.save(path)

    result = extract_text(path, ".pdf")

    assert (result.title, result.author) == ("Q3 Report", "Asha Verma")
    assert result.document_date == date(2024, 1, 15)


def test_metadata_drops_invisible_and_direction_changing_characters() -> None:
    # A right-to-left override can make "txt.exe" display as "exe.txt"; zero-width spaces hide text.
    assert extraction._clean_field("Invoice\u202eexe.pdf\u200b") == "Invoice exe.pdf"
    assert extraction._clean_field("\ufffd\x00  ") is None
    assert (
        extraction._clean_field("R\u00e9sum\u00e9, na\u00efve caf\u00e9")
        == "R\u00e9sum\u00e9, na\u00efve caf\u00e9"
    )
    assert extraction._clean_field("x" * 500) == "x" * 300


def test_pdf_without_metadata_has_none(tmp_path: Path) -> None:
    result = extract_text(make_pdf(tmp_path / "plain.pdf", ["Text."]), ".pdf")

    assert (result.title, result.author, result.document_date) == (None, None, None)


def test_docx_properties_title_author_date_and_tables(tmp_path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Leave Policy 2026"
    document.core_properties.author = "HR Team"
    document.core_properties.created = datetime(2026, 3, 1, 9, 30)
    document.add_paragraph("Intro.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "A"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "B"
    path = tmp_path / "policy.docx"
    document.save(path)

    result = extract_text(path, ".docx")

    assert (result.title, result.author, result.document_date) == (
        "Leave Policy 2026",
        "HR Team",
        date(2026, 3, 1),
    )
    assert result.table_count == 2


def test_docx_without_a_title_property_uses_its_first_heading(tmp_path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = ""
    document.add_heading("Onboarding Guide", level=1)
    document.add_paragraph("Welcome.")
    path = tmp_path / "guide.docx"
    document.save(path)

    assert extract_text(path, ".docx").title == "Onboarding Guide"


def test_markdown_title_is_the_first_top_level_heading(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("Intro line.\n\n## Details\n\n# Transformers\n\n# Later heading\n", encoding="utf-8")

    result = extract_text(path, ".md")

    assert result.title == "Transformers"
    assert result.author is None


def test_extraction_failures_carry_a_code(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.7 not a pdf")
    secret = make_pdf(
        tmp_path / "secret.pdf", ["x"], encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u"
    )

    for path, extension, code in (
        (broken, ".pdf", "damaged_file"),
        (secret, ".pdf", "password_protected"),
        (tmp_path / "gone.txt", ".txt", "file_missing"),
    ):
        with pytest.raises(ExtractionError) as caught:
            extract_text(path, extension)
        assert caught.value.code == code
