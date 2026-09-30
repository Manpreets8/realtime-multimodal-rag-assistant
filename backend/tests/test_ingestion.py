import uuid
from pathlib import Path

import docx
import pymupdf
import pytest
from httpx import AsyncClient
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models import DocumentChunk, DocumentStatus
from app.rag.embeddings import EmbeddingError
from app.services.ingestion_service import process_document
from tests.conftest import RegisterFn, bearer
from tests.fakes import HashingEmbeddingProvider, RecordingQueue

pytestmark = pytest.mark.integration

DOCS = "/api/v1/documents"


def pdf_bytes(pages: list[str], **save_options) -> bytes:
    document = pymupdf.open()
    for content in pages:
        page = document.new_page()
        if content:
            page.insert_textbox(pymupdf.Rect(72, 72, 540, 760), content, fontsize=10)
    return document.tobytes(**save_options)


def docx_bytes(tmp_path: Path) -> bytes:
    document = docx.Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees receive 18 days of annual paid leave each calendar year.")
    document.add_heading("Remote Work", level=1)
    document.add_paragraph("Remote work requires written approval from a manager.")
    path = tmp_path / "policy.docx"
    document.save(path)
    return path.read_bytes()


@pytest.fixture
async def alice(register_user: RegisterFn) -> dict[str, str]:
    return bearer((await register_user(email="alice@example.com"))["access_token"])


@pytest.fixture
async def kb_id(client: AsyncClient, alice: dict) -> str:
    return (await client.post("/api/v1/knowledge-bases", json={"name": "Handbook"}, headers=alice)).json()[
        "id"
    ]


async def upload(client: AsyncClient, headers: dict, kb: str, filename: str, content: bytes) -> dict:
    response = await client.post(
        f"{DOCS}/upload", data={"knowledge_base_id": kb}, files={"file": (filename, content)}, headers=headers
    )
    assert response.status_code == 201, response.text
    return response.json()


async def test_upload_enqueues_the_document(
    client: AsyncClient, alice: dict, kb_id: str, ingestion_queue: RecordingQueue
) -> None:
    document = await upload(client, alice, kb_id, "notes.txt", b"Some notes.")

    assert document["status"] == "uploaded"
    assert ingestion_queue.enqueued == [uuid.UUID(document["id"])]


async def test_pdf_is_extracted_chunked_embedded_and_stored(
    client: AsyncClient,
    alice: dict,
    kb_id: str,
    db: AsyncSession,
    embedder: HashingEmbeddingProvider,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "chunk_size", 300)
    monkeypatch.setattr(get_settings(), "chunk_overlap", 50)
    page_one = " ".join(f"Clause {i}: annual leave accrues at 1.5 days per month." for i in range(12))
    content = pdf_bytes([page_one, "Page two covers remote work approval by managers."])
    document = await upload(client, alice, kb_id, "handbook.pdf", content)

    status = await process_document(uuid.UUID(document["id"]))

    assert status is DocumentStatus.COMPLETED
    detail = (await client.get(f"{DOCS}/{document['id']}", headers=alice)).json()
    assert detail["status"] == "completed"
    assert detail["page_count"] == 2
    assert detail["chunk_count"] > 2
    assert detail["error_message"] is None and detail["error_code"] is None
    assert detail["processed_at"] is not None
    metadata = detail["extracted_metadata"]
    assert metadata["word_count"] == len(
        (page_one + " Page two covers remote work approval by managers.").split()
    )
    assert metadata["character_count"] > 0 and metadata["section_count"] == 0
    stats = detail["processing_stats"]
    stages = ("extraction_ms", "chunking_ms", "embedding_ms", "indexing_ms")
    assert all(isinstance(stats[stage], int) and stats[stage] >= 0 for stage in stages)
    assert stats["total_ms"] >= max(stats[stage] for stage in stages)
    assert stats["embedding_model"] == embedder.model_name

    chunks = (await client.get(f"{DOCS}/{document['id']}/chunks", headers=alice)).json()
    assert len(chunks) == detail["chunk_count"]
    assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert all(c["char_count"] <= 300 for c in chunks)
    assert chunks[-1]["page_number"] == 2
    assert "remote work" in chunks[-1]["content"].lower()
    # The embedder saw exactly the stored chunk texts.
    assert embedder.document_calls == [[c["content"] for c in chunks]]


async def test_vectors_and_full_text_index_are_queryable(
    client: AsyncClient,
    alice: dict,
    kb_id: str,
    db: AsyncSession,
    embedder: HashingEmbeddingProvider,
    tmp_path: Path,
) -> None:
    document = await upload(client, alice, kb_id, "policy.docx", docx_bytes(tmp_path))
    await process_document(uuid.UUID(document["id"]))

    query = await embedder.embed_query("How much annual paid leave do employees receive?")
    nearest = await db.scalar(
        select(DocumentChunk.content).order_by(DocumentChunk.embedding.cosine_distance(query)).limit(1)
    )
    keyword_hits = list(
        await db.scalars(
            select(DocumentChunk.section).where(
                DocumentChunk.content_tsv.op("@@")(func.plainto_tsquery("english", "approval manager"))
            )
        )
    )

    assert nearest is not None and "18 days" in nearest
    assert keyword_hits == ["Remote Work"]


@pytest.mark.parametrize(
    ("filename", "content", "message", "code"),
    [
        (
            "scan.pdf",
            pdf_bytes(["", ""]),
            "No readable text was found in this document. It may be a scanned",
            "no_text",
        ),
        (
            "secret.pdf",
            pdf_bytes(["Top secret"], encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u"),
            "password-protected",
            "password_protected",
        ),
    ],
)
async def test_unprocessable_documents_fail_with_a_reason(
    client: AsyncClient,
    alice: dict,
    kb_id: str,
    db: AsyncSession,
    filename: str,
    content: bytes,
    message: str,
    code: str,
) -> None:
    document = await upload(client, alice, kb_id, filename, content)

    status = await process_document(uuid.UUID(document["id"]))

    detail = (await client.get(f"{DOCS}/{document['id']}", headers=alice)).json()
    assert status is DocumentStatus.FAILED
    assert detail["status"] == "failed"
    assert message in detail["error_message"]
    assert detail["error_code"] == code
    assert detail["chunk_count"] == 0
    assert detail["processing_stats"] is None


async def test_embedding_failure_marks_document_failed_and_stores_no_chunks(
    client: AsyncClient, alice: dict, kb_id: str, db: AsyncSession, embedder: HashingEmbeddingProvider
) -> None:
    embedder.fail_with = EmbeddingError("The embedding service could not be reached.")
    document = await upload(client, alice, kb_id, "notes.md", b"# Notes\n\nImportant things.")

    await process_document(uuid.UUID(document["id"]))

    detail = (await client.get(f"{DOCS}/{document['id']}", headers=alice)).json()
    assert detail["status"] == "failed"
    assert detail["error_message"] == "The embedding service could not be reached."
    assert detail["error_code"] == "embedding_failed"  # temporary: worth retrying
    assert await db.scalar(select(func.count()).select_from(DocumentChunk)) == 0


async def test_unexpected_errors_are_not_leaked_to_users(
    client: AsyncClient, alice: dict, kb_id: str, db: AsyncSession, embedder: HashingEmbeddingProvider
) -> None:
    embedder.fail_with = RuntimeError("connection string postgres://secret@internal")
    document = await upload(client, alice, kb_id, "notes.txt", b"Plain notes.")

    await process_document(uuid.UUID(document["id"]))

    detail = (await client.get(f"{DOCS}/{document['id']}", headers=alice)).json()
    assert detail["status"] == "failed"
    assert "secret" not in detail["error_message"]
    assert "internal error" in detail["error_message"]
    assert detail["error_code"] == "internal_error"


async def test_reprocess_replaces_chunks_and_rejects_documents_in_flight(
    client: AsyncClient,
    alice: dict,
    kb_id: str,
    db: AsyncSession,
    embedder: HashingEmbeddingProvider,
    ingestion_queue: RecordingQueue,
) -> None:
    embedder.fail_with = EmbeddingError("temporary outage")
    document = await upload(client, alice, kb_id, "notes.txt", b"Leave policy: 18 days per year.")
    document_id = uuid.UUID(document["id"])
    await process_document(document_id)

    embedder.fail_with = None
    response = await client.post(f"{DOCS}/{document_id}/reprocess", headers=alice)
    assert response.status_code == 202
    assert response.json()["status"] == "uploaded"
    assert response.json()["error_message"] is None and response.json()["error_code"] is None
    assert ingestion_queue.enqueued[-1] == document_id

    again = await client.post(f"{DOCS}/{document_id}/reprocess", headers=alice)
    assert again.status_code == 409  # already queued

    await process_document(document_id)
    await client.post(f"{DOCS}/{document_id}/reprocess", headers=alice)
    await process_document(document_id)  # re-indexing a completed document
    count = await db.scalar(select(func.count()).where(DocumentChunk.document_id == document_id))
    assert count == 1


async def test_chunks_endpoint_is_owner_only(
    client: AsyncClient, alice: dict, kb_id: str, register_user: RegisterFn
) -> None:
    document = await upload(client, alice, kb_id, "notes.txt", b"Private notes.")
    await process_document(uuid.UUID(document["id"]))
    bob = bearer((await register_user(email="bob@example.com"))["access_token"])

    assert (await client.get(f"{DOCS}/{document['id']}/chunks", headers=bob)).status_code == 404
    assert (await client.post(f"{DOCS}/{document['id']}/reprocess", headers=bob)).status_code == 404


async def test_deleting_a_document_deletes_its_chunks(
    client: AsyncClient, alice: dict, kb_id: str, db: AsyncSession
) -> None:
    document = await upload(client, alice, kb_id, "notes.txt", b"Some notes to index.")
    await process_document(uuid.UUID(document["id"]))

    await client.delete(f"{DOCS}/{document['id']}", headers=alice)

    assert await db.scalar(select(func.count()).select_from(DocumentChunk)) == 0


async def test_processing_a_deleted_document_is_a_no_op(db: AsyncSession) -> None:
    assert await process_document(uuid.uuid4()) is None


# --- queue ---------------------------------------------------------------------
