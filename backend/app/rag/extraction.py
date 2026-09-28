"""Text extraction from PDF, DOCX, TXT and Markdown files.

Output is a list of `TextBlock`s: one per PDF page, or one per heading section
for DOCX/Markdown. Blocks are the unit the chunker never crosses, so every chunk
maps to exactly one page or section for citations.

All functions here are blocking (file I/O + parsing); call them from a worker thread.
"""

import codecs
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.rag.text_cleaning import clean_text

MAX_PDF_PAGES = 2000
# Zip-bomb guards for DOCX (a DOCX is a zip archive).
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
MAX_SECTION_LENGTH = 300


class ExtractionError(Exception):
    """Extraction failed for a reason the user can act on. The message is shown to them."""


@dataclass(frozen=True, slots=True)
class TextBlock:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    blocks: list[TextBlock]
    page_count: int | None = None


def extract_text(path: Path, extension: str) -> ExtractedDocument:
    if not path.is_file():
        raise ExtractionError("The uploaded file is missing from storage. Please upload it again.")
    match extension.lower():
        case ".pdf":
            return _extract_pdf(path)
        case ".docx":
            return _extract_docx(path)
        case ".md" | ".markdown":
            return _extract_markdown(path)
        case ".txt":
            return ExtractedDocument(blocks=_non_empty([TextBlock(clean_text(_read_text(path)))]))
        case _:
            raise ExtractionError(f"No text extractor is available for '{extension}' files.")


# --- PDF ----------------------------------------------------------------------


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        document = pymupdf.open(path, filetype="pdf")
    except (pymupdf.FileDataError, pymupdf.EmptyFileError, RuntimeError) as exc:
        raise ExtractionError("The PDF is damaged or could not be opened.") from exc

    with document:
        if document.needs_pass:
            raise ExtractionError("The PDF is password-protected. Remove the password and upload it again.")
        if document.page_count > MAX_PDF_PAGES:
            raise ExtractionError(f"The PDF has more than {MAX_PDF_PAGES} pages.")

        blocks: list[TextBlock] = []
        for page in document:
            # "blocks" keeps paragraph structure; sort=True gives natural reading order.
            paragraphs = [
                _join_block_lines(block[4])
                for block in page.get_text("blocks", sort=True)
                if block[6] == 0  # 0 = text block, 1 = image block
            ]
            text = clean_text("\n\n".join(paragraphs))
            blocks.append(TextBlock(text=text, page_number=page.number + 1))
        return ExtractedDocument(blocks=_non_empty(blocks), page_count=document.page_count)


_HYPHEN_LINE_BREAK = re.compile(r"(\w)-\n(\w)")


def _join_block_lines(text: str) -> str:
    """A PDF text block is one paragraph wrapped over several lines: rejoin words
    hyphenated at a line end first, then join the remaining lines with spaces."""
    return _HYPHEN_LINE_BREAK.sub(r"\1\2", text).replace("\n", " ")


# --- DOCX ---------------------------------------------------------------------


def _check_zip_safety(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise ExtractionError("The Word document is damaged or could not be opened.") from exc
    uncompressed = sum(entry.file_size for entry in entries)
    compressed = sum(entry.compress_size for entry in entries) or 1
    if uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES or uncompressed / compressed > MAX_DOCX_COMPRESSION_RATIO:
        raise ExtractionError("The Word document is too large when decompressed and was not processed.")


def _heading_level(paragraph: Paragraph) -> int | None:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style == "Title":
        return 1
    match = re.fullmatch(r"Heading (\d)", style)
    return int(match.group(1)) if match else None


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            # Merged cells are repeated by python-docx; keep one copy.
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_docx(path: Path) -> ExtractedDocument:
    _check_zip_safety(path)
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises a variety of parser errors
        raise ExtractionError("The Word document is damaged or could not be opened.") from exc

    sections = _SectionBuilder()
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item)
            if level is not None:
                sections.heading(level, text)
            else:
                sections.add(text)
        elif isinstance(item, Table):
            sections.add(_table_text(item))
    return ExtractedDocument(blocks=sections.blocks())


# --- Markdown / text ----------------------------------------------------------

_ATX_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE = re.compile(r"^\s*(```|~~~)")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return raw.decode("utf-16")
    return raw.decode("utf-8-sig")


def _extract_markdown(path: Path) -> ExtractedDocument:
    sections = _SectionBuilder()
    in_code = False
    for line in _read_text(path).splitlines():
        if _FENCE.match(line):
            in_code = not in_code
        heading = None if in_code else _ATX_HEADING.match(line)
        if heading:
            sections.heading(len(heading.group(1)), heading.group(2))
        else:
            sections.add(line, separator="\n")
    return ExtractedDocument(blocks=sections.blocks())


class _SectionBuilder:
    """Group content under its heading path ("Policies > Leave > Annual leave")."""

    def __init__(self) -> None:
        self._stack: list[tuple[int, str]] = []
        self._current: list[str] = []
        self._current_section: str | None = None
        self._blocks: list[TextBlock] = []

    def heading(self, level: int, title: str) -> None:
        self._flush()
        title = " ".join(title.split())
        self._stack = [(lvl, text) for lvl, text in self._stack if lvl < level] + [(level, title)]
        self._current_section = " > ".join(text for _, text in self._stack)[:MAX_SECTION_LENGTH]
        # The heading itself is part of the section's text: it helps retrieval.
        self._current.append(title)

    def add(self, text: str, separator: str = "\n\n") -> None:
        if self._current and separator == "\n\n":
            self._current.append("")
        self._current.append(text)

    def _flush(self) -> None:
        text = clean_text("\n".join(self._current))
        if text:
            self._blocks.append(TextBlock(text=text, section=self._current_section))
        self._current = []

    def blocks(self) -> list[TextBlock]:
        self._flush()
        return self._blocks


def _non_empty(blocks: list[TextBlock]) -> list[TextBlock]:
    return [block for block in blocks if block.text.strip()]
