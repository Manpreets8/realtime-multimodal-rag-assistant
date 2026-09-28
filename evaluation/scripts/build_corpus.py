"""Regenerate the binary evaluation documents (PDF and DOCX) from the text below.

The Markdown and text documents in ../corpus are edited directly; the PDF and DOCX are
built here so their content is reviewable and reproducible. Run with the backend venv:

    python evaluation/scripts/build_corpus.py
"""

from pathlib import Path

import docx
import pymupdf

CORPUS = Path(__file__).resolve().parents[1] / "corpus"

HANDBOOK_PAGES: list[list[tuple[str, str]]] = [
    [
        ("Northwind Labs Employee Handbook", ""),
        (
            "Working hours",
            "Core collaboration hours are 10:00 to 16:00 local time. Outside core hours, employees may "
            "arrange their schedule with their team. The standard work week is 40 hours for full-time "
            "employees.",
        ),
        (
            "Probation",
            "The probation period is 90 days. Performance is reviewed with your manager at 45 and 90 days.",
        ),
    ],
    [
        (
            "Annual leave",
            "Full-time employees receive 20 days of paid annual leave per calendar year. Part-time "
            "employees receive leave pro-rated to their contracted hours. Up to 5 unused days may be "
            "carried over to the next year and must be used by 31 March; any other unused leave expires. "
            "Leave requests should be submitted in the HR portal at least two weeks in advance for "
            "absences longer than three days.",
        ),
        (
            "Public holidays",
            "Northwind observes 11 public holidays per year. The list is published every December.",
        ),
    ],
    [
        (
            "Sick leave",
            "Employees receive 10 days of paid sick leave per year. A doctor's note is required for "
            "absences of more than three consecutive working days. Sick leave does not carry over to the "
            "next year.",
        ),
        (
            "Parental leave",
            "The primary caregiver receives 16 weeks of fully paid parental leave; the secondary caregiver "
            "receives 4 weeks. Parental leave must begin within 12 months of the birth or adoption.",
        ),
    ],
    [
        (
            "Remote work",
            "Employees may work remotely up to three days per week with their manager's approval. Remote "
            "work is not available during the first 90 days of employment (the probation period). "
            "Employees working abroad for more than 30 days in a year need prior approval from HR "
            "because of tax rules.",
        ),
    ],
    [
        (
            "Code of conduct",
            "Treat colleagues, customers and partners with respect. Concerns can be reported to the People "
            "team or anonymously through the EthicsLine at ethics.northwind.example. Retaliation against "
            "anyone who reports a concern in good faith is prohibited.",
        ),
    ],
]

BENEFITS: list[tuple[str, str]] = [
    (
        "Health insurance",
        "Health, dental and vision coverage starts on the first day of the month after your start date. "
        "Northwind pays 90% of the premium for employees and 75% for dependants.",
    ),
    (
        "Retirement plan",
        "Northwind matches retirement contributions up to 4% of salary. Matching contributions vest "
        "immediately.",
    ),
    (
        "Learning budget",
        "Each employee has an annual learning budget of $1,500 for courses, books and conferences. "
        "Unused budget does not roll over.",
    ),
    (
        "Home office",
        "New employees receive a one-time home office stipend of $500 for equipment such as a desk or chair.",
    ),
    (
        "Wellness",
        "A wellness allowance of $50 per month can be used for gym memberships, fitness classes or mental "
        "health apps.",
    ),
    (
        "Employee Assistance Programme",
        "The Employee Assistance Programme offers free, confidential counselling 24 hours a day.",
    ),
]


def build_pdf(path: Path) -> None:
    document = pymupdf.open()
    for sections in HANDBOOK_PAGES:
        page = document.new_page()
        y = 72
        for heading, body in sections:
            page.insert_text((72, y), heading, fontsize=14 if body else 18, fontname="helv")
            y += 26
            if body:
                rect = pymupdf.Rect(72, y, 540, y + 200)
                page.insert_textbox(rect, body, fontsize=11, fontname="helv")
                y += 140
    document.set_metadata({"title": "Northwind Labs Employee Handbook"})
    document.save(path, deflate=True)


def build_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Northwind Labs Benefits Guide", level=0)
    for heading, body in BENEFITS:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.save(path)


if __name__ == "__main__":
    build_pdf(CORPUS / "employee_handbook.pdf")
    build_docx(CORPUS / "benefits_guide.docx")
    print(f"Wrote {CORPUS / 'employee_handbook.pdf'} and {CORPUS / 'benefits_guide.docx'}")
